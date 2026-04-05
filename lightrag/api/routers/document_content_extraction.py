from __future__ import annotations

from functools import lru_cache
from io import BytesIO
from pathlib import Path


@lru_cache(maxsize=1)
def _is_docling_available() -> bool:
    """Check if docling is available (cached check)."""
    try:
        import docling  # noqa: F401  # type: ignore[import-not-found]

        return True
    except ImportError:
        return False


def _convert_with_docling(file_path: Path) -> str:
    """Convert document using docling (synchronous)."""
    from docling.document_converter import DocumentConverter  # type: ignore

    converter = DocumentConverter()
    result = converter.convert(file_path)
    return result.document.export_to_markdown()


def _extract_pdf_pypdf(file_bytes: bytes, password: str = None) -> str:
    """Extract PDF content using pypdf (synchronous)."""
    from pypdf import PdfReader  # type: ignore

    pdf_file = BytesIO(file_bytes)
    reader = PdfReader(pdf_file)

    if reader.is_encrypted:
        if not password:
            raise Exception("PDF is encrypted but no password provided")

        decrypt_result = reader.decrypt(password)
        if decrypt_result == 0:
            raise Exception("Incorrect PDF password")

    content = ""
    for page in reader.pages:
        content += page.extract_text() + "\n"

    return content


def _escape_tabular_cell(cell_value: str | int | float | None) -> str:
    """Escape characters that would break tab-delimited layout."""
    if cell_value is None:
        return ""
    text = str(cell_value)
    return (
        text.replace("\\", "\\\\")
        .replace("\t", "\\t")
        .replace("\r\n", "\\n")
        .replace("\r", "\\n")
        .replace("\n", "\\n")
    )


def _sanitize_sheet_title(title: str) -> str:
    """Escape sheet title to prevent formatting issues in separators."""
    return str(title).replace("\n", " ").replace("\t", " ").replace("\r", " ")


def _extract_docx(file_bytes: bytes) -> str:
    """Extract DOCX content including tables in document order (synchronous)."""
    from docx import Document  # type: ignore
    from docx.table import Table  # type: ignore
    from docx.text.paragraph import Paragraph  # type: ignore

    docx_file = BytesIO(file_bytes)
    doc = Document(docx_file)
    content_parts = []
    in_table = False

    for element in doc.element.body:
        if element.tag.endswith("p"):
            if in_table:
                content_parts.append("")
                in_table = False

            paragraph = Paragraph(element, doc)
            content_parts.append(paragraph.text)

        elif element.tag.endswith("tbl"):
            if content_parts and not in_table:
                content_parts.append("")

            in_table = True
            table = Table(element, doc)
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    row_text.append(_escape_tabular_cell(cell.text))
                if any(cell for cell in row_text):
                    content_parts.append("\t".join(row_text))

    return "\n".join(content_parts)


def _extract_pptx(file_bytes: bytes) -> str:
    """Extract PPTX content (synchronous)."""
    from pptx import Presentation  # type: ignore

    pptx_file = BytesIO(file_bytes)
    prs = Presentation(pptx_file)
    content = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                content += shape.text + "\n"
    return content


def _extract_xlsx(file_bytes: bytes) -> str:
    """Extract XLSX content in tab-delimited format with clear sheet separation."""
    from openpyxl import load_workbook  # type: ignore

    xlsx_file = BytesIO(file_bytes)
    wb = load_workbook(xlsx_file)
    content_parts: list[str] = []
    sheet_separator = "=" * 20

    for idx, sheet in enumerate(wb):
        if idx > 0:
            content_parts.append("")

        safe_title = _sanitize_sheet_title(sheet.title)
        content_parts.append(f"{sheet_separator} Sheet: {safe_title} {sheet_separator}")

        max_columns = sheet.max_column if sheet.max_column else 0

        for row in sheet.iter_rows(values_only=True):
            row_parts = []

            for idx in range(max_columns):
                if idx < len(row):
                    row_parts.append(_escape_tabular_cell(row[idx]))
                else:
                    row_parts.append("")

            if all(part == "" for part in row_parts):
                content_parts.append("")
            else:
                content_parts.append("\t".join(row_parts))

    content_parts.append(sheet_separator)
    return "\n".join(content_parts)
